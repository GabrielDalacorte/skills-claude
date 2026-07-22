# -*- coding: utf-8 -*-
"""Gera um DOCX-guia (capa + seções + prints + legendas) a partir de um doc.json.

Uso:  python gerar_docx.py <doc.json> [pastaDasImagens]

doc.json (todos os campos são opcionais menos title e output):
{
  "title": "Visibilidade",
  "subtitle": "Como funciona ...\\nsegunda linha",
  "dateline": "Guia com telas reais — Julho 2026",
  "intro": { "heading": "O que é", "paragraphs": ["..."], "bullets": ["..."] },
  "toc": ["1. ...", "2. ..."],
  "sections": [
     { "level": 1, "heading": "1. ...",
       "paragraphs": ["..."],
       "bullets": ["..."],           // lista com marcadores
       "steps": ["..."],             // lista numerada
       "figures": [ { "img": "01.png", "caption": "Figura 1 — ..." } ],
       "page_break_before": true      // default: true para level 1 (menos o 1º)
     },
     { "level": 2, "heading": "1.1 ...", "paragraphs": [...], "figures": [...] }
  ],
  "glossary": [ ["Termo", "definição"], ... ],
  "output": "D:\\\\AFL\\\\presentations\\\\X_Como_Funciona.docx",
  "image_width_in": 6.3            // largura das imagens (default 6.3in ~ 16cm em A4)
}
Ordem de blocos por seção: paragraphs -> bullets -> steps -> figures.
"""
import json
import os
import sys
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

spec = json.load(open(sys.argv[1], encoding='utf-8'))
IMGDIR = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(sys.argv[1]))
IMG_W = Inches(spec.get('image_width_in', 6.3))

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2)
    s.left_margin = Cm(2.5); s.right_margin = Cm(2.5)


def figura(img, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(IMGDIR, img), width=IMG_W)
    if caption:
        c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption); r.italic = True
        r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def bloco(sec):
    for t in sec.get('paragraphs', []):
        doc.add_paragraph(t)
    for b in sec.get('bullets', []):
        doc.add_paragraph(b, style='List Bullet')
    for i, s in enumerate(sec.get('steps', []), 1):
        doc.add_paragraph(f'{i}. {s}')
    for f in sec.get('figures', []):
        figura(f['img'], f.get('caption', ''))


# ── Capa ──────────────────────────────────────────────────────────────────
doc.add_paragraph(); doc.add_paragraph()
t = doc.add_heading(spec['title'], level=0); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
if spec.get('subtitle'):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(spec['subtitle']); r.font.size = Pt(15); r.font.color.rgb = RGBColor(0x4A, 0x4A, 0x4A)
if spec.get('dateline'):
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(spec['dateline']); r.font.size = Pt(11); r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
doc.add_page_break()

# ── Introdução + Sumário ────────────────────────────────────────────────────
intro = spec.get('intro')
if intro:
    doc.add_heading(intro.get('heading', 'Visão geral'), level=1)
    bloco(intro)
if spec.get('toc'):
    doc.add_heading('Sumário', level=2)
    for it in spec['toc']:
        doc.add_paragraph(it)
if intro or spec.get('toc'):
    doc.add_page_break()

# ── Seções ──────────────────────────────────────────────────────────────────
first_l1 = True
for sec in spec.get('sections', []):
    lvl = sec.get('level', 1)
    pbb = sec.get('page_break_before', lvl == 1 and not first_l1)
    if pbb:
        doc.add_page_break()
    if lvl == 1:
        first_l1 = False
    doc.add_heading(sec['heading'], level=lvl)
    bloco(sec)

# ── Glossário ────────────────────────────────────────────────────────────────
if spec.get('glossary'):
    doc.add_page_break()
    doc.add_heading('Glossário', level=1)
    for termo, desc in spec['glossary']:
        p = doc.add_paragraph()
        r = p.add_run(f'{termo}: '); r.bold = True
        p.add_run(desc)

doc.save(spec['output'])
print('DOCX salvo em:', spec['output'])
